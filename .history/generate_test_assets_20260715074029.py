import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment
from datetime import datetime, timedelta

# Create workbook
wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'Assets'

# Headers
headers = ['Asset Name', 'Asset Tag', 'Serial Number', 'Category', 'Purchase Cost', 'Purchase Date', 'Branch', 'Assigned To', 'Condition', 'Useful Life (months)']
ws.append(headers)

# Style header row
header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
header_font = Font(bold=True, color='FFFFFF')
for cell in ws[1]:
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal='center', vertical='center')

# Sample data
categories = ['Machinery', 'Vehicles', 'Furniture', 'Equipment', 'Technology', 'Tools', 'Building', 'Fixtures']
conditions = ['New', 'Good', 'Fair', 'Poor']
branches = ['HQ', 'Lagos', 'Abuja', 'Kano', 'Port Harcourt']

# Add 30 assets
for i in range(1, 31):
    cat_idx = (i - 1) % len(categories)
    cond_idx = (i - 1) % len(conditions)
    branch_idx = (i - 1) % len(branches)
    
    asset_name = f'{categories[cat_idx]} - Unit {i:02d}'
    asset_tag = f'AST-{i:04d}'
    serial = f'SN{i:06d}'
    category = categories[cat_idx]
    cost = 50000 + (i * 5000)
    date = (datetime.now() - timedelta(days=i*30)).strftime('%d/%m/%Y')
    branch = branches[branch_idx]
    assigned = f'staff{i % 5 + 1}@company.com'
    condition = conditions[cond_idx]
    life = 12 + (i % 36)
    
    ws.append([asset_name, asset_tag, serial, category, cost, date, branch, assigned, condition, life])

# Adjust column widths
ws.column_dimensions['A'].width = 25
ws.column_dimensions['B'].width = 12
ws.column_dimensions['C'].width = 12
ws.column_dimensions['D'].width = 12
ws.column_dimensions['E'].width = 15
ws.column_dimensions['F'].width = 15
ws.column_dimensions['G'].width = 15
ws.column_dimensions['H'].width = 20
ws.column_dimensions['I'].width = 12
ws.column_dimensions['J'].width = 18

# Save
wb.save('test-assets-30.xlsx')
print('✓ Generated test-assets-30.xlsx with 30 sample assets')
print('  File contains assets across 5 branches with realistic data')
print('  Use this file to test the asset import modal')
