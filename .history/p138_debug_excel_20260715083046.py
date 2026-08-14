#!/usr/bin/env python3
"""
Debug: See what's actually in the Excel files we're generating.
"""
import tempfile
from openpyxl import Workbook

# Test case that's failing
headers = ['name', 'assetTag', 'status', 'condition']
rows = [['Laptop 1', 'AST-001', 'active', 'good']]

wb = Workbook()
ws = wb.active
ws.append(headers)
for row in rows:
    ws.append(row)

tmp = tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False)
wb.save(tmp.name)

# Read it back to see what's there
from openpyxl import load_workbook
wb2 = load_workbook(tmp.name)
ws2 = wb2.active

print("Workbook contents:")
for row_idx, row in enumerate(ws2.iter_rows(min_row=1, max_row=2, values_only=False), 1):
    print(f"  Row {row_idx}:")
    for cell in row:
        print(f"    {cell.coordinate}: value={cell.value!r}, type={type(cell.value).__name__}, data_type={cell.data_type}")

print("\nValues only:")
for row_idx, row in enumerate(ws2.iter_rows(min_row=1, max_row=2, values_only=True), 1):
    print(f"  Row {row_idx}: {row}")
